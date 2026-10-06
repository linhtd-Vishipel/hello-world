"""Read an uploaded spreadsheet (.xlsx/.xlsm/.xls/.csv) into clean tables
and infer what each column holds, so the frontend can build a dashboard
without the uploader describing the data.

Output per sheet:
    {"name", "columns": [{"name", "type", "role", ...stats}], "rows": [[...]],
     "skipped_total_rows": int, "truncated": bool}

Column types: number | date | category | text.
Column roles: measure (sum/avg in KPIs and charts), dimension (group-by),
time (x axis), identifier (STT/Mã — excluded from KPIs), text (free text).
"""
import csv
import io
import re
from datetime import date, datetime, time, timedelta

MAX_ROWS_PER_SHEET = 200_000
MAX_HEADER_SCAN = 20

ALLOWED_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".csv"}

_TOTAL_ROW_RE = re.compile(r"^\s*(tổng|tong|total|cộng|cong|sum)\b", re.IGNORECASE)
_ID_NAME_RE = re.compile(r"^(stt|tt|no\.?|id|mã|ma|code|số hiệu|so hieu)\b|\b(id|mã|code)$", re.IGNORECASE)
_YEAR_NAME_RE = re.compile(r"^(năm|nam|year|tháng|thang|month|quý|quy|quarter)$", re.IGNORECASE)
_NUM_STRIP_RE = re.compile(r"[\s ]|(vnd|vnđ|đ|usd|\$|€)$", re.IGNORECASE)
_THOUSANDS_DOT_RE = re.compile(r"^-?\d{1,3}(\.\d{3})+$")
_THOUSANDS_COMMA_RE = re.compile(r"^-?\d{1,3}(,\d{3})+$")
_DATE_PATTERNS = [
    (re.compile(r"^\d{4}-\d{1,2}-\d{1,2}([ T]\d{1,2}:\d{2}(:\d{2})?)?$"), ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M")),
    (re.compile(r"^\d{1,2}/\d{1,2}/\d{4}( \d{1,2}:\d{2}(:\d{2})?)?$"), ("%d/%m/%Y", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S")),
    (re.compile(r"^\d{1,2}-\d{1,2}-\d{4}$"), ("%d-%m-%Y",)),
    (re.compile(r"^\d{1,2}\.\d{1,2}\.\d{4}$"), ("%d.%m.%Y",)),
]


class SpreadsheetError(ValueError):
    """Raised for files that cannot be read as a spreadsheet."""


# ---------------------------------------------------------------------------
# Value parsing
# ---------------------------------------------------------------------------


def parse_number(value):
    """Return a float for numeric-looking values (incl. Vietnamese
    formatting such as "1.234.567", "1.234,5", "12%", "500.000 đ"), else None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    s = _NUM_STRIP_RE.sub("", value.strip())
    if not s:
        return None
    pct = s.endswith("%")
    if pct:
        s = s[:-1]
    neg = s.startswith("(") and s.endswith(")")
    if neg:
        s = s[1:-1]
    if "." in s and "," in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif _THOUSANDS_DOT_RE.match(s):
        s = s.replace(".", "")
    elif _THOUSANDS_COMMA_RE.match(s):
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return None
    if n != n or n in (float("inf"), float("-inf")):
        return None
    if neg:
        n = -n
    return n / 100 if pct else n


def parse_date(value):
    """Return a datetime for date-like values, else None."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if not isinstance(value, str):
        return None
    s = value.strip()
    for pattern, formats in _DATE_PATTERNS:
        if pattern.match(s):
            for fmt in formats:
                try:
                    return datetime.strptime(s, fmt)
                except ValueError:
                    continue
    return None


def _is_empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _clean_cell(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, timedelta):
        return str(v)
    if isinstance(v, float) and v.is_integer() and abs(v) < 2**53:
        return int(v)
    return v


# ---------------------------------------------------------------------------
# Readers -> list of (sheet_name, rows-as-lists)
# ---------------------------------------------------------------------------


def _read_xlsx(content: bytes):
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for bad files
        raise SpreadsheetError(f"Không đọc được file Excel: {exc}") from exc
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_ROWS_PER_SHEET + MAX_HEADER_SCAN:
                    break
            yield ws.title, rows
    finally:
        wb.close()


def _read_xls(content: bytes):
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise SpreadsheetError(f"Không đọc được file Excel (.xls): {exc}") from exc
    for sheet in book.sheets():
        rows = []
        for r in range(min(sheet.nrows, MAX_ROWS_PER_SHEET + MAX_HEADER_SCAN + 1)):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:
                        row.append(cell.value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        yield sheet.name, rows


def _read_csv(content: bytes):
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        rows.append(row)
        if len(rows) > MAX_ROWS_PER_SHEET + MAX_HEADER_SCAN:
            break
    yield "CSV", rows


# ---------------------------------------------------------------------------
# Table extraction & type inference
# ---------------------------------------------------------------------------


def _find_header_row(rows):
    """Index of the first row that looks like a header: mostly non-empty
    text cells, spanning at least half of the widest row nearby."""
    scan = rows[:MAX_HEADER_SCAN]
    width = max((sum(not _is_empty(v) for v in r) for r in scan), default=0)
    if width == 0:
        return None
    for i, r in enumerate(scan):
        filled = [v for v in r if not _is_empty(v)]
        texts = [v for v in filled if isinstance(v, str) and parse_number(v) is None]
        if len(filled) >= max(1, width // 2) and len(texts) >= max(1, len(filled) * 0.6):
            return i
    # No text header: fall back to the first non-empty row as data with generic names.
    return -1


def _unique_names(raw):
    names, seen = [], {}
    for i, v in enumerate(raw):
        base = str(v).strip() if not _is_empty(v) else f"Cột {i + 1}"
        base = re.sub(r"\s+", " ", base)
        n = seen.get(base, 0)
        seen[base] = n + 1
        names.append(base if n == 0 else f"{base} ({n + 1})")
    return names


def _infer_column(name, values):
    non_empty = [v for v in values if v is not None]
    total = len(non_empty)
    info = {"name": name, "non_empty": total, "empty": len(values) - total}
    if total == 0:
        info.update(type="text", role="text", distinct=0)
        return info, values

    nums = [parse_number(v) for v in non_empty]
    num_ok = sum(n is not None for n in nums)
    dates = [parse_date(v) for v in non_empty]
    date_ok = sum(d is not None for d in dates)
    distinct = len({str(v) for v in non_empty})
    info["distinct"] = distinct

    if date_ok >= 0.8 * total and date_ok >= num_ok:
        out = []
        for v in values:
            d = parse_date(v) if v is not None else None
            out.append(d.isoformat(timespec="seconds") if d else None)
        ds = [d for d in out if d]
        info.update(type="date", role="time", min=min(ds), max=max(ds))
        return info, out

    if num_ok >= 0.8 * total:
        out = [parse_number(v) if v is not None else None for v in values]
        out = [int(n) if n is not None and n.is_integer() and abs(n) < 2**53 else n for n in out]
        ns = [n for n in out if n is not None]
        all_int = all(isinstance(n, int) for n in ns)
        info.update(type="number", min=min(ns), max=max(ns), sum=sum(ns), integer=all_int)
        if _YEAR_NAME_RE.match(name) and all_int and distinct <= 60:
            info.update(type="category", role="dimension")
            out = [str(n) if n is not None else None for n in out]
        elif _ID_NAME_RE.search(name) or (all_int and distinct == len(ns) and len(ns) > 20 and _is_sequence(ns)):
            info["role"] = "identifier"
        else:
            info["role"] = "measure"
        return info, out

    out = [str(v) if v is not None else None for v in values]
    if _ID_NAME_RE.search(name) and distinct > 0.9 * total:
        info.update(type="text", role="identifier")
    elif distinct <= max(30, 0.5 * total) and distinct < 1000:
        info.update(type="category", role="dimension")
    else:
        info.update(type="text", role="text")
    return info, out


def _is_sequence(ns):
    s = sorted(ns)
    return s[-1] - s[0] == len(s) - 1


def extract_table(sheet_name, raw_rows):
    rows = [[_clean_cell(v) for v in r] for r in raw_rows]
    rows = [r for r in rows if any(not _is_empty(v) for v in r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]

    header_idx = _find_header_row(rows)
    if header_idx is None:
        return None
    if header_idx >= 0:
        header, body = rows[header_idx], rows[header_idx + 1:]
    else:
        header, body = [None] * width, rows

    # Drop columns that are empty in both header and body.
    keep = [c for c in range(width) if not _is_empty(header[c]) or any(not _is_empty(r[c]) for r in body)]
    header = [header[c] for c in keep]
    body = [[r[c] for c in keep] for r in body]

    # Drop summary rows ("Tổng cộng") so totals are not double-counted.
    skipped_total = 0
    kept_body = []
    for r in body:
        first_text = next((v for v in r if isinstance(v, str)), None)
        if first_text is not None and _TOTAL_ROW_RE.match(first_text):
            skipped_total += 1
            continue
        kept_body.append(r)
    body = kept_body

    truncated = len(body) > MAX_ROWS_PER_SHEET
    body = body[:MAX_ROWS_PER_SHEET]
    if not body:
        return None

    names = _unique_names(header)
    columns, col_values = [], []
    for c, name in enumerate(names):
        info, vals = _infer_column(name, [r[c] for r in body])
        columns.append(info)
        col_values.append(vals)
    data = [list(r) for r in zip(*col_values)]
    return {
        "name": sheet_name,
        "columns": columns,
        "rows": data,
        "row_count": len(data),
        "skipped_total_rows": skipped_total,
        "truncated": truncated,
    }


def parse_spreadsheet(filename: str, content: bytes):
    """Parse an uploaded file into a list of sheet tables."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise SpreadsheetError("Chỉ hỗ trợ file .xlsx, .xlsm, .xls, .csv")
    reader = {".xlsx": _read_xlsx, ".xlsm": _read_xlsx, ".xls": _read_xls, ".csv": _read_csv}[ext]
    sheets = []
    for name, raw in reader(content):
        table = extract_table(name, raw)
        if table:
            sheets.append(table)
    if not sheets:
        raise SpreadsheetError("File không có dữ liệu dạng bảng")
    return sheets
